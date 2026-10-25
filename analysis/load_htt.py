"""
Adapter for the heat tolerance test (HTT) dataset: an EXTERNAL test set for
Ember, never used to train or calibrate anything.

    "HTT data set.xlsx", Mendeley Data, DOI 10.17632/g5sb382smp.1
    Licence CC BY 4.0. Study methods: PMC10464765 (San Diego, 2019).

65 adults (32 women, 33 men; the anthropometric sheets list 33 + 33, one
woman withdrew), each walking 120 min at 5.31 km/h, 4 % grade, once in the
heat (40 C, 40 % RH) and once thermoneutral (22 C, 40 % RH), in random
order. Core temperature and heart rate recorded every minute; a trial
stopped early at core >= 39 C, HR >= 180 bpm, or on request.

How it differs from PROSPIE (and why the numbers are not interchangeable):
  - Core reference: CorTemp INGESTIBLE PILL (method code "TP"), not rectal.
    Four trials say "RP" (undocumented; probably a rectal probe), eight
    are blank. All kept; the code is carried in `temp_method`.
  - Heart-rate sensor during the treadmill trials is not stated (the
    paper names a Polar H7 for the VO2max visit only).
  - No seated rest at the start: walking begins at minute 0. The resting
    heart rate is the one measured at rest and listed on the
    anthropometric sheets (`rest_hr`), which plays the role of BASELINE.
  - Ambient temperature is logged in Fahrenheit; converted to C here.

Cleaning rules, fixed before any result was looked at:
  1. Core outside 34-42 C -> missing (the same rule as load_prospie).
     Catches pill dropouts such as 58.5, 48.4 and 17.0 C.
  2. Trials whose notes say the pill was still in the stomach -> core
     missing for the whole trial (neutral 48 and 54): drinking moves a
     stomach pill by degrees.
  3. Pill artifacts. (a) A single reading DROP_C or more above both
     neighbours -> missing. (b) An ingestible pill reads a drink of water
     as a sudden fall in "core" that the body never had, then a slow climb
     back: a fall of DROP_C or more below the median of the last three
     valid readings (within DROP_LOOKBACK_MIN minutes) starts an artifact,
     and core is missing until it is back within RECOVER_C of that level,
     for at most RECOVER_MAX_MIN minutes. True core cannot jump or fall
     0.3 C in a minute while walking. clean_dips=False skips rule 3
     (reported as a sensitivity check).
These rules were set after looking at the raw traces and notes, before any
model was run on this data. Heart rate outside 30-230 bpm -> missing.
Missing minutes stay in the frame so the filter coasts through them,
exactly as with PROSPIE.
"""

import os
import sys

import numpy as np
import pandas as pd

DEFAULT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "HTT data set.xlsx")

SHEETS = {"heat": "Treadmill tests - HEAT",
          "neutral": "Treadmill tests - THERMO NEUTRA"}
DEMO_SHEETS = ("Anthropometric & Vo2max WOMEN", "Anthropometric & Vo2max MEN")
N_MIN = 121                       # minutes 0..120
COL_T0 = 2                        # column of minute 0
COL_METHOD, COL_AMB, COL_RH, COL_NOTES = 123, 126, 127, 129

DROP_C = 0.3
DROP_LOOKBACK_MIN = 10
RECOVER_C = 0.1
RECOVER_MAX_MIN = 30
PILL_IN_STOMACH = {("neutral", 48), ("neutral", 54)}


def _num(seq):
    return np.array([v if isinstance(v, (int, float)) and not isinstance(v, bool)
                     else np.nan for v in seq], float)


def _demographics(wb):
    out = {}
    for name in DEMO_SHEETS:
        rows = list(wb[name].iter_rows(values_only=True))
        header = [str(h).strip() if h is not None else "" for h in rows[1]]
        i_sex, i_rest = header.index("Gender"), header.index("Resting HR")
        for r in rows[2:]:
            if isinstance(r[0], (int, float)):
                out[int(r[0])] = {"sex": str(r[i_sex]).strip().upper()[:1],
                                  "rest_hr": float(r[i_rest])
                                  if isinstance(r[i_rest], (int, float)) else np.nan}
    return out


def _spikes(core):
    """Single readings DROP_C or more above both valid neighbours (rule 3a)."""
    idx = np.where(np.isfinite(core))[0]
    mask = np.zeros(len(core), bool)
    for a, k, b in zip(idx[:-2], idx[1:-1], idx[2:]):
        if core[k] - core[a] >= DROP_C and core[k] - core[b] >= DROP_C:
            mask[k] = True
    return mask


def _dips(core):
    """
    Boolean mask of pill artifacts (rule 3). core: NaN = missing.
    First single upward spikes, so a spike is never taken as the level a
    "fall" has to recover to; then falls, measured from the median of the
    last three valid readings.
    """
    n = len(core)
    mask = _spikes(core)
    good = []                        # indices of valid, unmasked readings
    i = 0
    while i < n:
        if not np.isfinite(core[i]) or mask[i]:
            i += 1
            continue
        recent = [k for k in good[-3:] if i - k <= DROP_LOOKBACK_MIN]
        level = float(np.median(core[recent])) if recent else np.nan
        if np.isfinite(level) and level - core[i] >= DROP_C:
            j = i
            while (j < n and j - i < RECOVER_MAX_MIN
                   and not (np.isfinite(core[j]) and core[j] >= level - RECOVER_C)):
                mask[j] = True
                j += 1
            i = j
            continue
        good.append(i)
        i += 1
    return mask


def load_htt(path=DEFAULT_PATH, clean_dips=True, verbose=True):
    """
    Returns the standard frame (subject, sex, condition, minute, hr, core_c,
    hr_type) plus rest_hr, ambient_c, ambient_rh, temp_method, notes.
    Subjects are prefixed "HTT" so they never collide with PROSPIE ids.
    """
    if not os.path.exists(path):
        sys.exit(f"no such file: {path}\n"
                 "Download it from DOI 10.17632/g5sb382smp.1 (Mendeley Data)")
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    demo = _demographics(wb)

    frames, log = [], {"range": 0, "stomach": 0, "dips": 0, "hr_range": 0}
    for cond, sheet in SHEETS.items():
        rows = list(wb[sheet].iter_rows(values_only=True))
        minutes = list(rows[2][COL_T0:COL_T0 + N_MIN])
        if minutes != list(range(N_MIN)):
            sys.exit(f"{sheet}: expected minute headers 0..120, got {minutes[:5]}...")
        trials, cur = {}, None
        for r in rows[3:]:
            if r[0] is not None:
                cur = int(r[0])
            if r[1] in ("Temp", "Heart rate"):
                trials.setdefault(cur, {})[r[1]] = r

        for sid, d in trials.items():
            if "Temp" not in d or "Heart rate" not in d:
                continue
            t_row = d["Temp"]
            core = _num(t_row[COL_T0:COL_T0 + N_MIN])
            hr = _num(d["Heart rate"][COL_T0:COL_T0 + N_MIN])
            have = np.isfinite(core) | np.isfinite(hr)
            if not have.any():
                continue
            n = int(np.where(have)[0].max()) + 1        # trim the empty tail
            core, hr = core[:n], hr[:n]

            bad = np.isfinite(core) & ~((core >= 34.0) & (core <= 42.0))
            log["range"] += int(bad.sum())
            core[bad] = np.nan
            if (cond, sid) in PILL_IN_STOMACH:
                log["stomach"] += int(np.isfinite(core).sum())
                core[:] = np.nan
            if clean_dips:
                dip = _dips(core)
                log["dips"] += int(dip.sum())
                core[dip] = np.nan
            bad_hr = np.isfinite(hr) & ~((hr >= 30) & (hr <= 230))
            log["hr_range"] += int(bad_hr.sum())
            hr[bad_hr] = np.nan

            amb = t_row[COL_AMB]
            info = demo.get(sid, {"sex": None, "rest_hr": np.nan})
            frames.append(pd.DataFrame({
                "subject": f"HTT{sid:02d}",
                "sex": info["sex"] if info["sex"] in ("M", "F") else np.nan,
                "condition": cond,
                "minute": np.arange(n, dtype=float),
                "hr": hr,
                "core_c": core,
                "hr_type": np.nan,
                "rest_hr": info["rest_hr"],
                "ambient_c": (float(amb) - 32.0) * 5.0 / 9.0
                if isinstance(amb, (int, float)) else np.nan,
                "ambient_rh": float(t_row[COL_RH])
                if isinstance(t_row[COL_RH], (int, float)) else np.nan,
                "temp_method": t_row[COL_METHOD] or "blank",
                "notes": t_row[COL_NOTES] or "",
            }))

    df = pd.concat(frames, ignore_index=True)
    # Trials with no usable core (stomach pill) or no heart rate add nothing.
    keep = df.groupby(["subject", "condition"]).transform(
        lambda s: s.notna().any())[["hr", "core_c"]].all(axis=1)
    n_trials_all = df.groupby(["subject", "condition"]).ngroups
    df = df[keep].sort_values(["subject", "condition", "minute"]).reset_index(drop=True)
    df["trial_hr_type"] = np.nan

    if verbose:
        trials = df.groupby(["subject", "condition"])
        print(f"loaded {os.path.basename(path)}")
        print(f"  {df.subject.nunique()} subjects, {trials.ngroups} trials "
              f"({n_trials_all - trials.ngroups} dropped: no usable core or HR)")
        print(f"  sex: {df.dropna(subset=['sex']).groupby('sex').subject.nunique().to_dict()}")
        print(f"  trials by condition: {trials.condition.first().value_counts().to_dict()}")
        print(f"  core method: {trials.temp_method.first().value_counts().to_dict()}")
        print(f"  core set missing: {log['range']} out of 34-42 C, "
              f"{log['stomach']} pill-in-stomach, {log['dips']} drinking dips"
              f"{'' if clean_dips else ' (dip rule OFF)'}; HR out of range {log['hr_range']}")
        print(f"  {len(df)} minutes, {int(df.core_c.notna().sum())} scored, "
              f"{int(df.hr.notna().sum())} with HR; "
              f"core {df.core_c.min():.2f}-{df.core_c.max():.2f} C")
        amb = trials.ambient_c.first()
        cond = trials.condition.first()
        print(f"  ambient median: heat {amb[cond == 'heat'].median():.1f} C, "
              f"neutral {amb[cond == 'neutral'].median():.1f} C; heat-sheet trials "
              f"logged under 30 C: {int(((cond == 'heat') & (amb < 30)).sum())}")
    return df


if __name__ == "__main__":
    load_htt()
